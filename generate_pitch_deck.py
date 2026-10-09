from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def apply_title_slide_formatting(slide):
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Project Name: Tagline"
    title.text_frame.paragraphs[0].font.size = Pt(54)
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = "The Problem we solve and a massive hook.\n\nDid you know X people suffer from Y every day?"
    subtitle.text_frame.paragraphs[0].font.size = Pt(28)

def add_slide(prs, title_text, content_bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 102, 204)
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    
    for i, bullet in enumerate(content_bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(24)

# Create presentation
prs = Presentation()

# Slide 1
title_slide_layout = prs.slide_layouts[0]
slide1 = prs.slides.add_slide(title_slide_layout)
apply_title_slide_formatting(slide1)

# Slide 2
add_slide(prs, "The Solution & Value Prop", [
    "Introduce your solution: 'We built [Project Name]...'",
    "Core Value 1: Extremely Fast",
    "Core Value 2: Highly Cost-Effective",
    "Core Value 3: Secure & Scalable",
    "Why Now: Perfect timing in the current tech landscape"
])

# Slide 3
add_slide(prs, "The Demo / How It Works", [
    "Step 1: Input - The user submits their issue.",
    "Step 2: Process - Our AI Engine analyzes the data.",
    "Step 3: Output - Real-time actionable insights are generated.",
    "(Insert GIF or Mockup of the app in action here)",
    "Scan the QR code to try it yourself!"
])

# Slide 4
add_slide(prs, "Market & Business Model", [
    "Target Audience: Local governments, agencies, engaged citizens.",
    "Business Model: SaaS subscription / Freemium tiers.",
    "Market Size: $X Billion TAM, $Y Million SOM.",
    "Competitor Advantage: We are faster and cheaper than existing options."
])

# Slide 5
add_slide(prs, "Tech Stack & Architecture", [
    "Frontend: React, Next.js, TailwindCSS",
    "Backend: Node.js, Express, PostgreSQL",
    "Cloud/AI: Google Cloud, Gemini 1.5 Pro",
    "(Insert Architecture Flowchart showing how the components talk)"
])

# Slide 6
add_slide(prs, "Team & Roadmap", [
    "The Team: [Name 1] - Lead, [Name 2] - Developer, [Name 3] - Design",
    "Next Week: Beta Launch",
    "Next Month: Official Mobile Release",
    "Next Year: Expand to 5 major cities",
    "The Ask: Join us in revolutionizing this industry."
])

output_filename = "Hackathon_Pitch_Deck.pptx"
prs.save(output_filename)
print(f"Successfully generated {output_filename}")
