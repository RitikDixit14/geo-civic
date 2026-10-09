from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_advanced_ppt():
    prs = Presentation()

    # Slide 1: Main Title Slide (Layout 0)
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]

    title1.text = "City Issue Reporter"
    title1.text_frame.paragraphs[0].font.bold = True
    title1.text_frame.paragraphs[0].font.color.rgb = RGBColor(41, 128, 185) # Blue theme

    subtitle1.text = "AI-Powered Platform for Urban Problem Resolution\nAutomated Triage | Geospatial Intelligence | Real-Time Dashboard"
    subtitle1.text_frame.paragraphs[0].font.size = Pt(20)

    # Content for the rest of the slides
    slides_data = [
        {
            "title": "Member 1: Citizen-Facing Application (Frontend)",
            "points": [
                ("Role: Lead Frontend Developer (Reporting Flow)", 0),
                ("Developed a responsive web application using React and Vite to ensure a frictionless experience for citizens.", 1),
                ("Seamless Data Capture:", 0),
                ("Integrated the HTML5 Geolocation API to automatically capture accurate latitude/longitude coordinates.", 1),
                ("Engineered a robust photo upload system utilizing FormData to stream images directly to the backend.", 1),
                ("User Interface & Tracking:", 0),
                ("Designed an intuitive 'My Reports' interface styled with modern Tailwind CSS.", 1),
                ("Allows citizens to track the real-time status of their submitted tickets (Open, In Progress, Resolved).", 1)
            ]
        },
        {
            "title": "Member 2: AI Service & Image Analysis",
            "points": [
                ("Role: Machine Learning & Python Engineer", 0),
                ("Built an independent, high-performance microservice using Python and FastAPI.", 1),
                ("Automated Issue Classification:", 0),
                ("Processes uploaded citizen photos to automatically categorize urban issues (e.g., Potholes, Water Leaks).", 1),
                ("Evaluates the image context to assign a baseline Severity Score (Low, Medium, High).", 1),
                ("Advanced Duplicate Prevention:", 0),
                ("Implemented perceptual image hashing (via ImageHash and Pillow) to detect visually identical photos.", 1),
                ("Provides a scalable foundation for plugging in production CNN models (MobileNet/ResNet).", 1)
            ]
        },
        {
            "title": "Member 3: Core Backend & Priority Engine",
            "points": [
                ("Role: Backend Developer & Database Architect", 0),
                ("Architected a secure Node.js/Express REST API serving both the citizen app and authority dashboard.", 1),
                ("Geospatial Database (PostGIS):", 0),
                ("Designed a PostgreSQL database leveraging PostGIS geometry data types.", 1),
                ("Engineered complex spatial queries (ST_Distance) to automatically cluster and merge duplicate reports within a 50-meter radius.", 1),
                ("Algorithmic Priority Engine:", 0),
                ("Developed an algorithm that calculates dynamic ticket priority scores.", 1),
                ("Priority scales based on AI severity, aggregate duplicate counts, and time elapsed.", 1)
            ]
        },
        {
            "title": "Member 4: Authority Command Center",
            "points": [
                ("Role: Dashboard & Data Visualization Engineer", 0),
                ("Created a real-time 'Command Center' for city authorities to triage and manage civic issues.", 1),
                ("Geographic Hotspots:", 0),
                ("Integrated React-Leaflet and Carto map tiles to visually display issue clusters and pinpoint ticket coordinates.", 1),
                ("Advanced Data Analytics:", 0),
                ("Developed dynamic, interactive BarCharts using Recharts to track category volumes and resolution progress.", 1),
                ("Priority Inbox System:", 0),
                ("Built a modernized triage table with color-coded severity badges, visual priority progress bars, and instant status updaters.", 1)
            ]
        }
    ]

    # Layout 1 is Title and Content
    for data in slides_data:
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        
        # Style Title
        title = slide.shapes.title
        title.text = data["title"]
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(41, 128, 185) # Blue theme
        title.text_frame.paragraphs[0].font.size = Pt(32)
        title.text_frame.paragraphs[0].font.bold = True
        
        # Style Content
        content = slide.placeholders[1]
        tf = content.text_frame
        tf.clear()
        
        for text, level in data["points"]:
            p = tf.add_paragraph()
            p.text = text
            p.level = level
            
            # Format main bullet points differently from sub-bullets
            if level == 0:
                p.font.bold = True
                p.font.size = Pt(22)
                p.font.color.rgb = RGBColor(44, 62, 80) # Dark gray
                p.space_before = Pt(14)
            else:
                p.font.size = Pt(18)
                p.font.color.rgb = RGBColor(52, 73, 94) # Lighter gray
                p.space_before = Pt(6)

    # Save over the old one
    output_path = r"C:\Users\DELL\Desktop\City_Issue_Reporter_Presentation.pptx"
    prs.save(output_path)
    print("Advanced presentation generated successfully!")

if __name__ == "__main__":
    create_advanced_ppt()
